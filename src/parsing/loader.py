from __future__ import annotations

from pathlib import Path

from docx import Document
from pypdf import PdfReader

SUPPORTED = {'.pdf', '.docx', '.txt'}


def extract_text(path: str | Path) -> str:
    p = Path(path)
    suffix = p.suffix.lower()
    if suffix not in SUPPORTED:
        raise ValueError('Unsupported resume format. Use PDF, DOCX or TXT.')
    if suffix == '.pdf':
        reader = PdfReader(str(p))
        return '\n'.join((page.extract_text() or '') for page in reader.pages)
    if suffix == '.docx':
        doc = Document(str(p))
        chunks = [para.text for para in doc.paragraphs]
        for table in doc.tables:
            for row in table.rows:
                chunks.append(' | '.join(cell.text for cell in row.cells))
        return '\n'.join(chunks)
    return p.read_text(encoding='utf-8', errors='ignore')
